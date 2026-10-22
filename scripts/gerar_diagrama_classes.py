#!/usr/bin/env python3
"""
Gera o diagrama de classes do Dr. Moo em .docx (com as figuras em PNG).

O modelo aqui descrito é extraído do código: as entidades são as interfaces de
`packages/db/src/tipos.ts` e o esquema de `schema.ts`; o motor é a classe
`Sessao` de `packages/engine`; os serviços são os módulos de `db`, `sync`,
`planilha` e `relatorio`. Quando o código mudar, este arquivo precisa mudar
junto — é um diagrama de verdade, não um desenho decorativo.

Uso: python3 scripts/gerar_diagrama_classes.py
"""
from __future__ import annotations

import sys
from dataclasses import dataclass, field
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt

RAIZ = Path(__file__).resolve().parent.parent
DOCS = RAIZ / "docs"
IMAGENS = DOCS / "img"
SAIDA = DOCS / "DrMoo-DiagramaDeClasses.docx"

sys.path.insert(0, str(RAIZ / "scripts"))
from gerar_docx import (  # noqa: E402  (precisa do sys.path acima)
    CINZA,
    CINZA_CLARO,
    FONTE,
    LOGO,
    TEAL,
    converter,
    preparar_estilos,
    rodape,
)

# Paleta (mesma da logo, em hex para o matplotlib)
C_TEAL = "#015B66"
C_TEAL_CLARO = "#0E7F8C"
C_FUNDO = "#FFFFFF"
C_BORDA = "#C9D4D7"
C_TEXTO = "#0F1F2C"
C_CINZA = "#424E57"
C_VERDE = "#2E8B5A"
C_AMBAR = "#B26A00"

LARGURA_CARACTERE = 0.078  # em polegadas, para a fonte usada nas caixas
ALTURA_LINHA = 0.185
ALTURA_TITULO = 0.40


@dataclass
class Classe:
    """Uma caixa do diagrama."""

    nome: str
    x: float
    y: float
    atributos: list[str] = field(default_factory=list)
    operacoes: list[str] = field(default_factory=list)
    estereotipo: str = ""
    cor: str = C_TEAL
    largura: float = 0.0

    def altura(self) -> float:
        linhas = len(self.atributos) + len(self.operacoes)
        separadores = 0.06 * (1 if self.operacoes else 0)
        return ALTURA_TITULO + linhas * ALTURA_LINHA + 0.14 + separadores

    def calcular_largura(self) -> float:
        if self.largura:
            return self.largura
        textos = [self.nome, self.estereotipo, *self.atributos, *self.operacoes]
        maior = max(len(t) for t in textos)
        return max(1.7, maior * LARGURA_CARACTERE + 0.28)


@dataclass
class Relacao:
    """Ligação entre duas caixas. `tipo` muda a ponta da seta."""

    de: str
    para: str
    rotulo: str = ""
    tipo: str = "associacao"  # associacao | composicao | heranca | dependencia
    lado_de: str = "auto"
    lado_para: str = "auto"


def ancora(classe: Classe, lado: str) -> tuple[float, float]:
    largura = classe.calcular_largura()
    altura = classe.altura()
    if lado == "topo":
        return classe.x + largura / 2, classe.y + altura
    if lado == "baixo":
        return classe.x + largura / 2, classe.y
    if lado == "esquerda":
        return classe.x, classe.y + altura / 2
    return classe.x + largura, classe.y + altura / 2


def lados_automaticos(a: Classe, b: Classe) -> tuple[str, str]:
    """Escolhe por onde sair e por onde entrar, pela posição relativa das caixas."""
    ax, ay = a.x + a.calcular_largura() / 2, a.y + a.altura() / 2
    bx, by = b.x + b.calcular_largura() / 2, b.y + b.altura() / 2
    if abs(bx - ax) >= abs(by - ay):
        return ("direita", "esquerda") if bx > ax else ("esquerda", "direita")
    return ("topo", "baixo") if by > ay else ("baixo", "topo")


def desenhar(classes: list[Classe], relacoes: list[Relacao], arquivo: Path, titulo: str) -> Path:
    por_nome = {c.nome: c for c in classes}

    largura_total = max(c.x + c.calcular_largura() for c in classes) + 0.35
    altura_total = max(c.y + c.altura() for c in classes) + 0.55

    fig, ax = plt.subplots(figsize=(largura_total, altura_total), dpi=200)
    ax.set_xlim(0, largura_total)
    ax.set_ylim(0, altura_total)
    ax.axis("off")
    fig.patch.set_facecolor(C_FUNDO)

    ax.text(
        0.1,
        altura_total - 0.22,
        titulo,
        fontsize=11,
        fontweight="bold",
        color=C_TEAL,
        va="center",
    )

    # Relações primeiro, para as caixas ficarem por cima das linhas.
    for rel in relacoes:
        a, b = por_nome[rel.de], por_nome[rel.para]
        lado_a, lado_b = lados_automaticos(a, b)
        if rel.lado_de != "auto":
            lado_a = rel.lado_de
        if rel.lado_para != "auto":
            lado_b = rel.lado_para
        xa, ya = ancora(a, lado_a)
        xb, yb = ancora(b, lado_b)

        estilos = {
            "composicao": dict(arrowstyle="-|>", linestyle="-", color=C_TEAL, lw=1.3),
            "heranca": dict(arrowstyle="-|>", linestyle="-", color=C_CINZA, lw=1.1),
            "dependencia": dict(arrowstyle="->", linestyle="--", color=C_CINZA, lw=1.0),
            "associacao": dict(arrowstyle="->", linestyle="-", color=C_CINZA, lw=1.0),
        }[rel.tipo]

        ax.add_patch(
            FancyArrowPatch(
                (xa, ya),
                (xb, yb),
                connectionstyle="arc3,rad=0.06",
                mutation_scale=11,
                shrinkA=2,
                shrinkB=2,
                zorder=1,
                **estilos,
            )
        )
        if rel.rotulo:
            # A 35% do caminho, e não no meio: no meio o rótulo encosta na
            # caixa de destino quando as caixas têm tamanhos diferentes.
            ax.text(
                xa + (xb - xa) * 0.35,
                ya + (yb - ya) * 0.35 + 0.08,
                rel.rotulo,
                fontsize=6.0,
                color=C_CINZA,
                ha="center",
                va="center",
                bbox=dict(boxstyle="round,pad=0.12", fc=C_FUNDO, ec="none"),
                zorder=2,
            )

    for c in classes:
        largura = c.calcular_largura()
        altura = c.altura()

        ax.add_patch(
            FancyBboxPatch(
                (c.x, c.y),
                largura,
                altura,
                boxstyle="round,pad=0,rounding_size=0.06",
                linewidth=1.1,
                edgecolor=C_BORDA,
                facecolor="#FFFFFF",
                zorder=3,
            )
        )
        # Faixa do título
        ax.add_patch(
            FancyBboxPatch(
                (c.x, c.y + altura - ALTURA_TITULO),
                largura,
                ALTURA_TITULO,
                boxstyle="round,pad=0,rounding_size=0.06",
                linewidth=0,
                facecolor=c.cor,
                zorder=4,
            )
        )
        topo = c.y + altura - ALTURA_TITULO
        if c.estereotipo:
            ax.text(
                c.x + largura / 2,
                topo + ALTURA_TITULO * 0.70,
                c.estereotipo,
                fontsize=5.6,
                color="#D9EEF0",
                ha="center",
                va="center",
                zorder=5,
            )
            ax.text(
                c.x + largura / 2,
                topo + ALTURA_TITULO * 0.30,
                c.nome,
                fontsize=7.6,
                fontweight="bold",
                color="#FFFFFF",
                ha="center",
                va="center",
                zorder=5,
            )
        else:
            ax.text(
                c.x + largura / 2,
                topo + ALTURA_TITULO / 2,
                c.nome,
                fontsize=7.8,
                fontweight="bold",
                color="#FFFFFF",
                ha="center",
                va="center",
                zorder=5,
            )

        y = topo - 0.09
        for atributo in c.atributos:
            ax.text(c.x + 0.10, y, atributo, fontsize=6.0, color=C_TEXTO, va="top", zorder=5)
            y -= ALTURA_LINHA

        if c.operacoes:
            y -= 0.02
            ax.plot([c.x, c.x + largura], [y + 0.04, y + 0.04], color=C_BORDA, lw=0.8, zorder=5)
            y -= 0.04
            for operacao in c.operacoes:
                ax.text(
                    c.x + 0.10, y, operacao, fontsize=6.0, color=C_TEAL_CLARO, va="top", zorder=5
                )
                y -= ALTURA_LINHA

    IMAGENS.mkdir(parents=True, exist_ok=True)
    fig.savefig(arquivo, bbox_inches="tight", facecolor=C_FUNDO)
    plt.close(fig)
    return arquivo


# ---------------------------------------------------------------------------
# Os três diagramas
# ---------------------------------------------------------------------------

def diagrama_dominio() -> tuple[Path, str]:
    """Entidades persistidas: o rebanho e o que acontece com ele."""
    classes = [
        Classe(
            "Propriedade",
            0.2,
            5.55,
            ["+ id: string", "+ nome: string", "+ areaPropriaHa: number", "+ precoLeiteLitro: number", "+ gestacaoDias: number"],
            estereotipo="«entidade»",
        ),
        Classe(
            "Lote",
            3.6,
            5.75,
            ["+ id: string", "+ nome: string", "+ descricao?: string"],
            estereotipo="«entidade»",
        ),
        Classe(
            "Animal",
            0.2,
            2.75,
            [
                "+ id: string",
                "+ brinco: string  «único»",
                "+ nome?: string",
                "+ categoria: CategoriaAnimal",
                "+ sexo: 'F' | 'M'",
                "+ raca?: string",
                "+ dataNascimento?: DataISO",
                "+ loteId?: string",
                "+ status: StatusAnimal",
                "+ paiId? / maeId?: string",
                "+ createdAt / updatedAt: DataISO",
                "+ deletedAt?: DataISO",
            ],
            estereotipo="«entidade raiz»",
        ),
        Classe(
            "AnimalComResumo",
            0.2,
            0.2,
            [
                "+ del?: number",
                "+ idadeMeses?: number",
                "+ producaoMediaDia?: number",
                "+ ultimoCmt?: ResultadoCMT",
                "+ cmtPositivoMeses: number",
                "+ carenciaLeiteAte?: DataISO",
                "+ dataProvavelParto?: DataISO",
            ],
            estereotipo="«projeção de leitura»",
            cor=C_TEAL_CLARO,
        ),
        Classe(
            "Lactacao",
            3.6,
            4.30,
            ["+ numero: number", "+ dataParto: DataISO", "+ dataSecagemReal?: DataISO", "+ encerrada: boolean"],
            estereotipo="«entidade»",
        ),
        Classe(
            "ControleLeiteiro",
            3.6,
            2.75,
            ["+ data: DataISO", "+ litrosManha: number", "+ litrosTarde: number", "+ litrosTotal: number"],
            estereotipo="«entidade»",
        ),
        Classe(
            "ExameCMT",
            3.6,
            1.45,
            ["+ data: DataISO", "+ ad / ae / pd / pe:", "   ResultadoCMT?"],
            estereotipo="«entidade»",
        ),
        Classe(
            "Atendimento",
            7.1,
            4.55,
            [
                "+ id: string",
                "+ dataHora: DataISO",
                "+ queixaPrincipal?: string",
                "+ kbVersao / engineVersao: string",
                "+ fatosIniciais: JSON",
                "+ respostas: JSON",
                "+ conclusoes: JSON",
                "+ traco: JSON",
                "+ status: 'rascunho' | 'concluido'",
            ],
            estereotipo="«entidade»",
            cor=C_VERDE,
        ),
        Classe(
            "Tratamento",
            7.1,
            2.45,
            [
                "+ dataInicio / dataFim: DataISO",
                "+ dose / via: string",
                "+ carenciaLeiteAte: DataISO",
                "+ custoTotal: number",
            ],
            estereotipo="«entidade»",
        ),
        Classe(
            "Medicamento",
            10.6,
            2.45,
            ["+ nome: string", "+ custoUnidade: number", "+ carenciaLeiteDias: number", "+ estoque: number"],
            estereotipo="«entidade»",
        ),
        Classe(
            "EventoCalendario",
            10.6,
            4.55,
            ["+ tipo: TipoEvento", "+ titulo: string", "+ data: DataISO", "+ status: string", "+ origem: string"],
            estereotipo="«entidade»",
        ),
        Classe(
            "Lancamento",
            7.1,
            0.20,
            ["+ data: DataISO", "+ categoria: Despesa|Receita", "+ valor: number", "+ animalId?: string"],
            estereotipo="«entidade»",
        ),
        Classe(
            "Cobertura",
            3.6,
            0.05,
            ["+ data: DataISO", "+ tipo: 'ia'|'monta'|'iatf'", "+ confirmadaPrenhe: boolean", "+ dataProvavelParto?: DataISO"],
            estereotipo="«entidade»",
        ),
    ]
    relacoes = [
        Relacao("Lote", "Animal", "0..1 agrupa *", "associacao"),
        Relacao("Animal", "Lactacao", "1 → *", "composicao"),
        Relacao("Animal", "ControleLeiteiro", "1 → *", "composicao"),
        Relacao("Animal", "ExameCMT", "1 → *", "composicao"),
        Relacao("Animal", "Atendimento", "1 → *", "composicao"),
        Relacao("Animal", "Tratamento", "1 → *", "composicao"),
        Relacao("Animal", "Cobertura", "1 → *", "composicao"),
        Relacao("Medicamento", "Tratamento", "0..1 usado em *", "associacao"),
        Relacao("Atendimento", "Tratamento", "origina", "dependencia"),
        Relacao("Atendimento", "EventoCalendario", "agenda", "dependencia"),
        Relacao("Animal", "Lancamento", "0..1 → *", "associacao"),
        Relacao("AnimalComResumo", "Animal", "estende", "heranca"),
        Relacao("Propriedade", "Lote", "1 → *", "composicao"),
    ]
    caminho = IMAGENS / "classes-dominio.png"
    desenhar(classes, relacoes, caminho, "Figura 1 — Domínio persistido (packages/db)")
    return caminho, "Domínio persistido"


def diagrama_motor() -> tuple[Path, str]:
    """O motor de inferência e a base de conhecimento."""
    classes = [
        Classe(
            "Sessao",
            0.2,
            2.3,
            [
                "- kb: BaseCompilada",
                "- fatosIniciais: Fato[]",
                "- respostas: RespostaUsuario[]",
                "- estado: EstadoDerivado",
            ],
            [
                "+ proximo(): ProximoPasso",
                "+ responder(r): void",
                "+ voltar(): boolean",
                "+ hipoteses(): Hipotese[]",
                "+ encerrar(): Parecer",
                "+ serializar(): Estado",
                "+ restaurar(kb, estado): Sessao",
            ],
            estereotipo="«serviço de domínio»",
            cor=C_VERDE,
        ),
        Classe(
            "BaseDeConhecimento",
            4.3,
            4.75,
            ["+ versao: string", "+ variavelObjetivo: string", "+ variaveis: VariavelDef[]", "+ hipoteses: HipoteseDef[]", "+ perguntas: Pergunta[]", "+ regras: Regra[]"],
            estereotipo="«JSON versionado»",
        ),
        Classe(
            "Regra",
            8.3,
            4.75,
            [
                "+ id: string",
                "+ grupo: GrupoRegra",
                "+ se: Condicao",
                "+ entao: Conclusao[]",
                "+ cf: CF",
                "+ status: 'validada'|'revisar'",
                "+ fontes: string[]",
            ],
            estereotipo="«regra de produção»",
        ),
        Classe(
            "Condicao",
            8.3,
            2.85,
            ["«simples» var op valor", "«todas» E", "«alguma» OU", "«nao» negação"],
            estereotipo="«união recursiva»",
            cor=C_AMBAR,
        ),
        Classe(
            "Fato",
            0.2,
            0.2,
            ["+ variavel: string", "+ valor: Valor", "+ cf: CF", "+ origem: OrigemFato", "+ em: DataISO"],
            estereotipo="«memória de trabalho»",
        ),
        Classe(
            "Pergunta",
            4.3,
            2.85,
            ["+ variavel: string", "+ tipo: TipoPergunta", "+ texto: string", "+ custo: number", "+ preRequisitos?: Condicao"],
            estereotipo="«entidade da base»",
        ),
        Classe(
            "Hipotese",
            4.3,
            0.2,
            ["+ id / rotulo: string", "+ cf: CF", "+ confianca: string", "+ bandeiraVermelha: boolean", "+ conduta: string[]", "+ regrasQueSustentam: string[]"],
            estereotipo="«saída»",
            cor=C_VERDE,
        ),
        Classe(
            "Parecer",
            8.3,
            0.2,
            ["+ motivoParada: MotivoParada", "+ hipoteses: Hipotese[]", "+ bandeiraVermelha?: Hipotese", "+ traco: PassoDoTraco[]", "+ kbVersao / engineVersao"],
            estereotipo="«saída»",
            cor=C_VERDE,
        ),
    ]
    relacoes = [
        Relacao("Sessao", "BaseDeConhecimento", "compila e usa", "associacao"),
        Relacao("BaseDeConhecimento", "Regra", "1 → 44", "composicao"),
        Relacao("BaseDeConhecimento", "Pergunta", "1 → *", "composicao"),
        Relacao("Regra", "Condicao", "se", "composicao"),
        Relacao("Pergunta", "Condicao", "preRequisitos", "dependencia"),
        Relacao("Sessao", "Fato", "mantém *", "composicao"),
        Relacao("Sessao", "Hipotese", "produz *", "dependencia"),
        Relacao("Sessao", "Parecer", "encerra com", "dependencia"),
        Relacao("Hipotese", "Parecer", "compõe", "associacao"),
    ]
    caminho = IMAGENS / "classes-motor.png"
    desenhar(classes, relacoes, caminho, "Figura 2 — Motor de inferência (packages/engine + knowledge)")
    return caminho, "Motor de inferência"


def diagrama_servicos() -> tuple[Path, str]:
    """Os serviços de aplicação: persistência, sincronização, planilha e relatório."""
    classes = [
        Classe(
            "Repositorios",
            0.2,
            4.00,
            ["«toda a SQL do app»"],
            [
                "+ listarRebanho(filtros)",
                "+ obterAnimal(id)",
                "+ salvarAnimal(a)",
                "+ painelMastite(ano, mes)",
                "+ panoramaDiagnosticos(dias)",
                "+ panoramaPorVaca(dias)",
                "+ resumoMensal(ano, mes)",
            ],
            estereotipo="«módulo @drmoo/db»",
            cor=C_TEAL,
        ),
        Classe(
            "Driver",
            0.2,
            1.95,
            ["- conexao: BackendSQLite", "- emMemoria: boolean"],
            ["+ abrirBanco()", "+ consultar / executar", "+ emTransacao(fn)", "+ definirBackend(b)"],
            estereotipo="«acesso ao SQLite»",
        ),
        Classe(
            "BackendSQLite",
            0.2,
            0.20,
            ["«interface»"],
            ["execAsync / runAsync", "getAllAsync / getFirstAsync", "withTransactionAsync"],
            estereotipo="«porta»",
            cor=C_AMBAR,
        ),
        Classe(
            "expo-sqlite",
            3.3,
            0.95,
            ["nativo no Android", "WebAssembly no navegador"],
            estereotipo="«adaptador»",
            cor=C_CINZA,
        ),
        Classe(
            "node:sqlite",
            3.3,
            0.05,
            ["usado pela suíte de testes"],
            estereotipo="«adaptador»",
            cor=C_CINZA,
        ),
        Classe(
            "Sincronizacao",
            6.6,
            4.00,
            ["«ida e volta única»"],
            [
                "+ instantaneoLocal()",
                "+ aplicarLinhas(diff)",
                "+ sincronizar(opcoes)",
                "+ verificarServidor(url)",
            ],
            estereotipo="«módulo @drmoo/db»",
        ),
        Classe(
            "Mesclagem",
            10.2,
            4.00,
            ["«last-write-wins»"],
            ["+ mesclar(a, b)", "+ diferenca(base, alvo)", "+ validarInstantaneo(x)", "+ contarLinhas(i)"],
            estereotipo="«módulo @drmoo/sync»",
            cor=C_VERDE,
        ),
        Classe(
            "ServidorSync",
            6.6,
            1.90,
            ["porta 8787 na LAN", "código de 6 dígitos"],
            ["GET /saude", "POST /sincronizar"],
            estereotipo="«processo Node»",
            cor=C_CINZA,
        ),
        Classe(
            "Planilha",
            3.3,
            3.30,
            ["ABAS: AbaPlanilha[]"],
            ["+ exportarPlanilha()", "+ importarPlanilha(bytes)", "+ abaParaBanco(...)", "+ bancoParaAba(...)"],
            estereotipo="«@drmoo/planilha + db»",
        ),
        Classe(
            "Relatorio",
            3.3,
            1.60,
            ["HTML → PDF"],
            ["+ htmlProntuario(d)", "+ htmlEstatisticas(d)", "+ graficoLinha / Barras / Rosca"],
            estereotipo="«módulo @drmoo/relatorio»",
        ),
        Classe(
            "Telas (Expo Router)",
            0.2,
            6.55,
            ["rebanho · saúde · anamnese", "prontuário · financeiro", "estatísticas · rede local"],
            estereotipo="«apps/mobile»",
            cor=C_TEAL_CLARO,
        ),
    ]
    relacoes = [
        Relacao("Telas (Expo Router)", "Repositorios", "consulta", "dependencia"),
        Relacao("Telas (Expo Router)", "Sincronizacao", "", "dependencia"),
        Relacao("Repositorios", "Driver", "usa", "composicao"),
        Relacao("Driver", "BackendSQLite", "depende da porta", "associacao"),
        Relacao("BackendSQLite", "expo-sqlite", "implementa", "heranca"),
        Relacao("BackendSQLite", "node:sqlite", "implementa", "heranca"),
        Relacao("Sincronizacao", "Mesclagem", "usa", "dependencia"),
        Relacao("Sincronizacao", "Driver", "", "dependencia"),
        Relacao("ServidorSync", "Mesclagem", "mesma função", "dependencia"),
        Relacao("Planilha", "Driver", "grava", "dependencia"),
        Relacao("Relatorio", "Repositorios", "lê números de", "dependencia"),
    ]
    caminho = IMAGENS / "classes-servicos.png"
    desenhar(classes, relacoes, caminho, "Figura 3 — Serviços de aplicação e portas")
    return caminho, "Serviços de aplicação"


# ---------------------------------------------------------------------------
# Documento
# ---------------------------------------------------------------------------

LEGENDA = """
## Como ler o diagrama

| Traço | Significado |
|---|---|
| Linha cheia com ponta fechada | composição: a parte não existe sem o todo (um exame pertence a um animal) |
| Linha cheia com ponta aberta | associação: uma classe conhece a outra |
| Linha tracejada | dependência: uma usa a outra sem guardar referência |
| Ponta fechada partindo de uma «porta» | realização: o adaptador implementa a interface |

Os estereótipos entre «guillemets» dizem o papel da classe: `«entidade»` é tabela no SQLite,
`«projeção de leitura»` é resultado de consulta (não existe no banco), `«porta»` é interface
que isola o resto do código de uma tecnologia, e `«adaptador»` é a implementação dessa porta.
"""

NOTAS = """
## Por que o modelo é assim

**O brinco é a identidade pública, o id é a interna.** Toda tabela usa `id TEXT PRIMARY KEY`
gerado pelo app, mas a planilha, a sincronização e a tela conversam pelo brinco do animal e pelo
nome do lote, do medicamento e da vacina — é o que o produtor conhece.

**`AnimalComResumo` não é tabela.** É o resultado de uma consulta única com subconsultas
correlacionadas (DEL, produção média, último CMT, carência, parto previsto). Agregar em SQL e não
em JavaScript é decisão de desempenho: o aparelho-alvo tem 2 GB de RAM.

**O atendimento guarda JSON de propósito.** `fatosIniciais`, `respostas`, `conclusoes` e `traco`
são a fotografia do raciocínio. Guardá-los permite **re-executar** o motor depois e provar que o
mesmo caso dá o mesmo parecer — e detectar quando a base mudou de versão.

**`BackendSQLite` é uma porta.** O driver não conhece `expo-sqlite`: conhece uma interface de
cinco métodos. É isso que deixa a suíte rodar contra o SQLite de verdade do Node, sem navegador
nem aparelho, e o que permitiu à aba secundária do navegador cair para um banco em memória sem
que nada acima percebesse.

**Exclusão é lógica onde a sincronização precisa.** `animal.deleted_at` existe para que a
exclusão se propague entre aparelhos: sem lápide, o registro apagado num aparelho voltaria do
outro na sincronização seguinte.

**A mesclagem é a mesma função nos dois lados.** O servidor da propriedade e o aparelho chamam
`mesclar()` do mesmo pacote. É isso que garante convergência: mesclar A com B e B com A dá o
mesmo resultado, e repetir não muda nada.
"""


def main() -> None:
    figuras = [diagrama_dominio(), diagrama_motor(), diagrama_servicos()]

    doc = Document()
    for secao in doc.sections:
        # Paisagem: diagrama em retrato fica ilegível.
        secao.page_width = Inches(11.69)
        secao.page_height = Inches(8.27)
        secao.left_margin = Inches(0.6)
        secao.right_margin = Inches(0.6)
        secao.top_margin = Inches(0.6)
        secao.bottom_margin = Inches(0.6)

    preparar_estilos(doc)

    # Capa curta — o documento é um anexo técnico, não a documentação inteira.
    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(60)
        p.add_run().add_picture(str(LOGO), width=Inches(1.5))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Dr. Moo — Diagrama de classes")
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = TEAL
    r.font.name = FONTE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Modelo de domínio, motor de inferência e serviços de aplicação")
    r.font.size = Pt(12)
    r.font.color.rgb = CINZA
    r.font.name = FONTE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Hackathon UFJF 2026 · gerado a partir do código por scripts/gerar_diagrama_classes.py")
    r.font.size = Pt(9)
    r.font.color.rgb = CINZA_CLARO
    r.font.name = FONTE

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for caminho, titulo in figuras:
        doc.add_heading(titulo, level=1)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(caminho), width=Inches(10.3))
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    converter(doc, LEGENDA.strip())
    converter(doc, NOTAS.strip())

    rodape(doc)
    doc.save(SAIDA)

    for caminho, _ in figuras:
        print(f"  + {caminho.relative_to(RAIZ)}")
    print(f"\n{SAIDA.relative_to(RAIZ)} ({SAIDA.stat().st_size / 1024:.0f} KB)")


if __name__ == "__main__":
    main()
